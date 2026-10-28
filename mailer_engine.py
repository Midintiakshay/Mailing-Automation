#!/usr/bin/env python3
"""
Mailer Engine for Bigrock / Roundcube Webmail Automation
Decoupled, thread-safe automation worker with live progress and status callbacks.
"""

import csv
import html
import os
import random
import re
import threading
import time
from pathlib import Path
from typing import Callable, Dict, List, Optional, Any

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.chrome.service import Service
from webdriver_manager.chrome import ChromeDriverManager

# ---------------------------------------------------------------------------
# SELECTORS  (Roundcube Elastic + Larry + classic skins)
# ---------------------------------------------------------------------------
COMPOSE_CSS = [
    "a.compose", "a#rcmbtn107", "a#rcmbtn106", "#rcmbtn107", "#composebtn",
    "a[href*='_action=compose']", "a.button.compose", "a[onclick*='compose']",
    "[data-action='compose']",
]
COMPOSE_XPATH = [
    "//a[normalize-space(translate(@title,'COMPSE','compse'))='compose']",
    "//a[contains(translate(@title,'COMPSE','compse'),'compose')]",
    "//*[self::a or self::button][contains(translate(normalize-space(.),"
    "'ABCDEFGHIJKLMNOPQRSTUVWXYZ','abcdefghijklmnopqrstuvwxyz'),'compose')]",
    "//*[self::a or self::button][contains(translate(normalize-space(.),"
    "'ABCDEFGHIJKLMNOPQRSTUVWXYZ','abcdefghijklmnopqrstuvwxyz'),'new message')]",
]

TO_CSS = [
    "#compose_to input.recipient-input", "#compose_to input[role='combobox']",
    "#compose_to textarea", "#compose_to input", "textarea#_to", "input#_to",
    "textarea[name='_to']", "input[name='_to']", "input#rcmcomposeto",
    "input[role='combobox'][id*='to']",
]
TO_XPATH = ["//label[contains(.,'To')]/following::input[1]"]

SUBJECT_CSS = [
    "input#compose-subject", "input[name='_subject']", "#compose_subject input",
    "input#_subject", "input[placeholder*='ubject']",
]
SUBJECT_XPATH = ["//label[contains(.,'Subject')]/following::input[1]"]

SEND_CSS = [
    "a.send", "button.send", "a#rcmbtn106", "button#rcmbtn106", "#rcmbtn106-clone",
    "button[data-action='send']", "a[onclick*='send']", "button.btn-primary.send",
]
SEND_XPATH = [
    "//a[contains(translate(@title,'SENDMSG','sendmsg'),'send')]",
    "//*[self::a or self::button][translate(normalize-space(.),"
    "'SEND','send')='send']",
    "//*[self::a or self::button][contains(translate(normalize-space(.),"
    "'ABCDEFGHIJKLMNOPQRSTUVWXYZ','abcdefghijklmnopqrstuvwxyz'),'send')]",
]


# ---------------------------------------------------------------------------
# PARSING HELPERS
# ---------------------------------------------------------------------------
def extract_field(record: Dict[str, str], *keys) -> str:
    """Find value matching candidate keys, handling BOM, underscores, spaces, and case."""
    # 1. Exact match (case-insensitive & stripped)
    for k in keys:
        norm = k.strip().lower()
        if norm in record and str(record[norm]).strip():
            return str(record[norm]).strip()
    # 2. Normalized match (remove underscores, dashes, spaces)
    norm_rec = {
        re.sub(r'[\s_\-\.]+', '', k.lower()): str(v).strip()
        for k, v in record.items() if v and str(v).strip()
    }
    for k in keys:
        norm_k = re.sub(r'[\s_\-\.]+', '', k.lower())
        if norm_k in norm_rec and norm_rec[norm_k]:
            return norm_rec[norm_k]
    return ""


NAME_CANDIDATES = [
    "author_name", "author name", "name", "author", "authors", "author_names",
    "full_name", "fullname", "first_name", "first author", "corresponding author",
    "recipient", "recipient_name", "contact_name", "contact", "scholar", "researcher"
]

EMAIL_CANDIDATES = [
    "email", "e-mail", "email id", "email_id", "mail", "mail id", "mail_id",
    "author_email", "author email", "contact_email", "recipient_email"
]

TITLE_CANDIDATES = [
    "paper_title", "paper title", "title", "article_title", "article title",
    "paper", "article", "manuscript_title", "manuscript title", "work_title"
]

TOPIC_CANDIDATES = [
    "topic", "subject", "journal", "category", "field", "research_area", "discipline"
]


def clean_recipient_row(name, email, paper_title="", topic="", row_num=None):
    email = (email or "").strip()
    if not email or "@" not in email:
        return None
    return {
        "name": (name or "").strip(),
        "email": email,
        "paper_title": (paper_title or "").strip(),
        "topic": (topic or "").strip(),
    }


def parse_recipients_file(file_path: Path) -> List[Dict[str, str]]:
    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    rows = []
    suffix = file_path.suffix.lower()

    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(file_path, read_only=True, data_only=True)
        ws = wb.active
        header = None
        for row_num, raw in enumerate(ws.iter_rows(values_only=True), start=1):
            values = ["" if c is None else str(c).strip() for c in raw]
            if header is None:
                header = [v.lstrip('\ufeff').strip().lower() for v in values]
                continue
            record = {k: v for k, v in zip(header, values) if k}
            email = extract_field(record, *EMAIL_CANDIDATES)
            name = extract_field(record, *NAME_CANDIDATES)
            title = extract_field(record, *TITLE_CANDIDATES)
            topic = extract_field(record, *TOPIC_CANDIDATES)

            row = clean_recipient_row(name, email, title, topic, row_num=row_num)
            if row:
                rows.append(row)
    else:
        with open(file_path, newline="", encoding="utf-8-sig", errors="replace") as f:
            reader = csv.DictReader(f)
            for row_num, raw_record in enumerate(reader, start=2):
                record = {
                    (k or "").lstrip('\ufeff').strip().lower(): (v or "").strip()
                    for k, v in raw_record.items() if k
                }
                email = extract_field(record, *EMAIL_CANDIDATES)
                name = extract_field(record, *NAME_CANDIDATES)
                title = extract_field(record, *TITLE_CANDIDATES)
                topic = extract_field(record, *TOPIC_CANDIDATES)

                row = clean_recipient_row(name, email, title, topic, row_num=row_num)
                if row:
                    rows.append(row)
    return rows


def text_to_template_dict(text: str) -> Dict[str, str]:
    raw = text.strip().replace("\r\n", "\n")
    blocks = re.split(r'\n\s*\n', raw)
    html_paragraphs = []
    plain_paragraphs = []

    for b in blocks:
        lines = [l.strip() for l in b.split('\n') if l.strip()]
        if not lines:
            continue
        p_html = "<br>".join(html.escape(l) for l in lines)
        html_paragraphs.append(
            f'<p style="margin: 0 0 12pt 0; line-height: 1.5;">{p_html}</p>'
        )
        plain_paragraphs.append("\n".join(lines))

    full_html = (
        '<div style="font-family: \'Times New Roman\', Times, serif; font-size: 12pt; color: #000000; line-height: 1.5;">\n'
        + '\n'.join(html_paragraphs)
        + '\n</div>'
    )
    full_plain = '\n\n'.join(plain_paragraphs)
    return {"html": full_html, "plain": full_plain}


def parse_template_docx(file_path: Path) -> Dict[str, str]:
    from docx import Document

    doc = Document(str(file_path))
    paras = list(doc.paragraphs)
    while paras and not paras[-1].text.strip():
        paras.pop()

    html_paragraphs = []
    plain_paragraphs = []

    for p in paras:
        p_html_runs = []
        p_plain_runs = []
        for r in p.runs:
            raw_text = r.text
            escaped = html.escape(raw_text).replace('\n', '<br>')
            if r.bold:
                escaped = f'<b>{escaped}</b>'
            if r.italic:
                escaped = f'<i>{escaped}</i>'
            if r.underline:
                escaped = f'<u>{escaped}</u>'

            styles = []
            if r.font.name:
                styles.append(f"font-family: '{r.font.name}', Times, serif")
            if r.font.size:
                styles.append(f"font-size: {r.font.size.pt}pt")
            if r.font.color and r.font.color.rgb:
                styles.append(f"color: #{r.font.color.rgb}")

            if styles:
                escaped = f'<span style="{"; ".join(styles)}">{escaped}</span>'

            p_html_runs.append(escaped)
            p_plain_runs.append(raw_text)

        html_content = ''.join(p_html_runs)
        plain_content = ''.join(p_plain_runs)

        if not html_content.strip() or html_content.strip() == '<br>':
            html_paragraphs.append('<p style="margin: 0 0 12pt 0; line-height: 1.5;"><br></p>')
            plain_paragraphs.append('')
        else:
            align_style = ""
            if p.paragraph_format.alignment:
                align_name = str(p.paragraph_format.alignment).split('.')[-1].lower()
                if align_name in ("center", "right", "justify"):
                    align_style = f" text-align: {align_name};"

            html_paragraphs.append(
                f'<p style="margin: 0 0 12pt 0; line-height: 1.5;{align_style}">{html_content}</p>'
            )
            plain_paragraphs.append(plain_content)

    full_html = (
        '<div style="font-family: \'Times New Roman\', Times, serif; font-size: 12pt; color: #000000; line-height: 1.5;">\n'
        + '\n'.join(html_paragraphs)
        + '\n</div>'
    )
    full_plain = '\n\n'.join(plain_paragraphs)
    full_plain = re.sub(r'\n{3,}', '\n\n', full_plain).strip()
    return {"html": full_html, "plain": full_plain}


def parse_template_file(file_path: Path) -> Dict[str, str]:
    if not file_path.exists():
        raise FileNotFoundError(f"Template not found: {file_path}")
    suffix = file_path.suffix.lower()
    if suffix in (".docx", ".docm", ".dotx", ".dotm"):
        return parse_template_docx(file_path)
    elif suffix in (".doc", ".dot", ".rtf", ".word"):
        # Legacy word if pywin32 available
        try:
            import win32com.client
            word = win32com.client.Dispatch("Word.Application")
            word.Visible = False
            word.DisplayAlerts = 0
            doc = word.Documents.Open(str(file_path), ReadOnly=True, AddToRecentFiles=False)
            try:
                text = doc.Range().Text
            finally:
                doc.Close(SaveChanges=False)
                word.Quit()
            return text_to_template_dict(text or "")
        except Exception:
            text = file_path.read_text(encoding="utf-8-sig", errors="replace")
            return text_to_template_dict(text)
    else:
        text = file_path.read_text(encoding="utf-8-sig", errors="replace")
        return text_to_template_dict(text)


def parse_subjects_file(file_path: Path) -> List[str]:
    if not file_path.exists():
        return []
    suffix = file_path.suffix.lower()
    raw_lines = []
    if suffix in (".docx", ".docm", ".dotx", ".dotm"):
        from docx import Document
        doc = Document(str(file_path))
        for p in doc.paragraphs:
            text = (p.text or "").strip()
            if text:
                for line in text.splitlines():
                    cleaned = line.strip()
                    if cleaned:
                        raw_lines.append(cleaned)
    else:
        text = file_path.read_text(encoding="utf-8-sig", errors="replace")
        for line in text.splitlines():
            cleaned = line.strip()
            if cleaned:
                raw_lines.append(cleaned)
    return raw_lines


def personalize_template(template_data: Dict[str, str], name: str, paper_title: str) -> Dict[str, str]:
    display_name = (name or "").strip()
    safe_name = html.escape(display_name) if display_name else "Researcher"
    fallback_name = display_name if display_name else "Researcher"
    paper_title_str = (paper_title or "").strip()
    safe_paper = html.escape(paper_title_str)

    html_out = template_data.get("html", "")
    plain_out = template_data.get("plain", "")

    # Replace author name placeholders (case-insensitive, handling any whitespace like { name })
    name_patterns = [
        r'\{\s*(name|author|author_name|authorname|recipient)\s*\}',
        r'\{name\}', r'\{NAME\}', r'\{Name\}',
        r'\{author_name\}', r'\{AUTHOR_NAME\}',
        r'\{author\}', r'\{AUTHOR\}',
    ]
    for pat in name_patterns:
        html_out = re.sub(pat, safe_name, html_out, flags=re.IGNORECASE)
        plain_out = re.sub(pat, fallback_name, plain_out, flags=re.IGNORECASE)

    # Replace paper title placeholders (case-insensitive, handling { title }, { paper_title }, etc.)
    title_patterns = [
        r'\{\s*(paper_title|papertitle|title|article_title)\s*\}',
        r'\{paper_title\}', r'\{PAPER_TITLE\}',
        r'\{title\}', r'\{TITLE\}',
    ]
    for pat in title_patterns:
        html_out = re.sub(pat, safe_paper, html_out, flags=re.IGNORECASE)
        plain_out = re.sub(pat, paper_title_str, plain_out, flags=re.IGNORECASE)

    # If the template literally contains "Dear Researcher", personalize it with the actual author name
    if display_name:
        html_out = re.sub(r'Dear\s+Researcher\b', f'Dear {safe_name}', html_out, flags=re.IGNORECASE)
        plain_out = re.sub(r'Dear\s+Researcher\b', f'Dear {fallback_name}', plain_out, flags=re.IGNORECASE)

    return {"html": html_out, "plain": plain_out}


def generate_fallback_subject(name: str, topic: str, idx: int) -> str:
    topic = (topic or "").strip()
    author = (name or "").strip()
    prefixes = (
        "Invitation to Submit",
        "Research Article Invitation",
        "Research Publication Opportunity",
        "Invitation for Manuscript Consideration",
        "Call for Research Submissions",
    )
    prefix = prefixes[idx % len(prefixes)]

    if topic and len(topic) <= 80:
        return f"{prefix}: {topic}"
    if author:
        return f"{prefix} for {author}"
    if topic:
        return f"{prefix} for Researcher"
    return prefix


# ---------------------------------------------------------------------------
# AUTOMATION ENGINE CLASS
# ---------------------------------------------------------------------------
class MailerEngine:
    """Manages browser automation, execution thread, pauses, and status updates."""

    def __init__(self):
        self.lock = threading.Lock()
        self.thread: Optional[threading.Thread] = None
        self.stop_requested = threading.Event()
        self.pause_requested = threading.Event()
        self.is_paused = False

        # Status tracking
        self.status = "idle"  # idle | running | paused | stopped | completed | error
        self.current_index = 0
        self.total_count = 0
        self.sent_count = 0
        self.failed_count = 0
        self.current_recipient = ""
        self.current_subject = ""
        self.error_message = ""
        self.start_time: Optional[float] = None

        # Callbacks
        self.log_callbacks: List[Callable[[str, str], None]] = []
        self.progress_callbacks: List[Callable[[Dict[str, Any]], None]] = []

        # Selenium driver reference
        self.driver = None

    def register_log_callback(self, cb: Callable[[str, str], None]):
        self.log_callbacks.append(cb)

    def register_progress_callback(self, cb: Callable[[Dict[str, Any]], None]):
        self.progress_callbacks.append(cb)

    def log(self, message: str, level: str = "info"):
        timestamp = time.strftime("%H:%M:%S")
        formatted = f"[{timestamp}] {message}"
        for cb in self.log_callbacks:
            try:
                cb(formatted, level)
            except Exception:
                pass

    def emit_progress(self, email: str = "", status_code: str = "", subject: str = "", error: str = ""):
        data = {
            "status": self.status,
            "current": self.current_index,
            "total": self.total_count,
            "sent": self.sent_count,
            "failed": self.failed_count,
            "current_recipient": self.current_recipient,
            "current_subject": self.current_subject,
            "email": email,
            "result_status": status_code,
            "subject": subject,
            "error": error,
            "timestamp": time.strftime("%Y-%m-%d %H:%M:%S"),
            "elapsed_seconds": int(time.time() - self.start_time) if self.start_time else 0,
        }
        for cb in self.progress_callbacks:
            try:
                cb(data)
            except Exception:
                pass

    def get_status_dict(self) -> Dict[str, Any]:
        with self.lock:
            return {
                "status": self.status,
                "current": self.current_index,
                "total": self.total_count,
                "sent": self.sent_count,
                "failed": self.failed_count,
                "current_recipient": self.current_recipient,
                "current_subject": self.current_subject,
                "error_message": self.error_message,
                "elapsed_seconds": int(time.time() - self.start_time) if self.start_time else 0,
            }

    def start(self, config: Dict[str, Any]):
        with self.lock:
            if self.thread and self.thread.is_alive():
                raise RuntimeError("Mailer is already running.")

            self.stop_requested.clear()
            self.pause_requested.clear()
            self.is_paused = False
            self.status = "running"
            self.sent_count = 0
            self.failed_count = 0
            self.current_index = 0
            self.error_message = ""
            self.start_time = time.time()

            self.thread = threading.Thread(
                target=self._run_job,
                args=(config,),
                daemon=True,
                name="MailerEngineWorker",
            )
            self.thread.start()

    def pause(self):
        with self.lock:
            if self.status == "running":
                self.pause_requested.set()
                self.is_paused = True
                self.status = "paused"
                self.log("Pause requested. Will pause before next email.", "warning")

    def resume(self):
        with self.lock:
            if self.status == "paused":
                self.pause_requested.clear()
                self.is_paused = False
                self.status = "running"
                self.log("Resumed sending process.", "info")

    def stop(self):
        with self.lock:
            if self.status in ("running", "paused"):
                self.stop_requested.set()
                self.pause_requested.clear()
                self.is_paused = False
                self.status = "stopped"
                self.log("Stop requested by user. Cleaning up...", "warning")

    # -----------------------------------------------------------------------
    # SELENIUM CORE AUTOMATION HELPERS
    # -----------------------------------------------------------------------
    def _start_browser(self, headless: bool = False):
        options = webdriver.ChromeOptions()
        if headless:
            options.add_argument("--headless=new")
            options.add_argument("--window-size=1920,1080")
        else:
            options.add_argument("--start-maximized")
        options.add_argument("--disable-popup-blocking")
        options.add_argument("--disable-notifications")
        options.add_experimental_option(
            "prefs", {"profile.default_content_setting_values.popups": 1}
        )
        service = Service(ChromeDriverManager().install())
        driver = webdriver.Chrome(service=service, options=options)
        if not headless:
            driver.maximize_window()
        return driver

    def _visible_one(self, css_list=(), xpath_list=(), timeout=10):
        deadline = time.time() + timeout
        last_err = None
        while time.time() < deadline:
            if self.stop_requested.is_set():
                raise InterruptedError("Process stopped by user.")
            for by, selectors in ((By.CSS_SELECTOR, css_list), (By.XPATH, xpath_list)):
                for sel in selectors:
                    try:
                        for el in self.driver.find_elements(by, sel):
                            if el.is_displayed() and el.is_enabled():
                                return el
                    except Exception as exc:
                        last_err = exc
            time.sleep(0.3)
        raise Exception(f"No visible element found for selectors ({last_err})")

    def _any_visible(self, css_list=(), xpath_list=()):
        try:
            self._visible_one(css_list, xpath_list, timeout=0.6)
            return True
        except Exception:
            return False

    def _safe_click(self, el):
        self.driver.execute_script("arguments[0].scrollIntoView({block:'center'});", el)
        try:
            el.click()
        except Exception:
            self.driver.execute_script("arguments[0].click();", el)

    def _dismiss_any_dialogs(self):
        try:
            buttons = self.driver.find_elements(
                By.CSS_SELECTOR,
                ".modal-dialog button, .modal-content button, #messagestack a, "
                "#confirm-dialog button, .ui-dialog-buttonset button, div.dialog button",
            )
            for btn in buttons:
                if btn.is_displayed() and any(
                    k in (btn.text or "").lower() for k in ("discard", "yes", "ok", "close")
                ):
                    self._safe_click(btn)
                    time.sleep(0.5)
                    break
        except Exception:
            pass

    def _dump_error(self, tag: str):
        if not self.driver:
            return
        stamp = int(time.time())
        try:
            self.driver.save_screenshot(f"error_{tag}_{stamp}.png")
        except Exception:
            pass
        try:
            Path(f"error_{tag}_{stamp}.html").write_text(self.driver.page_source, encoding="utf-8")
        except Exception:
            pass

    def _login(self, wait, webmail_url: str, from_email: str, password: str):
        self.log(f"Opening webmail login page: {webmail_url}...")
        self.driver.get(webmail_url)

        user = wait.until(
            EC.presence_of_element_located((By.CSS_SELECTOR, "#rcmloginuser, input[name='_user']"))
        )
        user.clear()
        user.send_keys(from_email)

        pwd = self.driver.find_element(By.CSS_SELECTOR, "#rcmloginpwd, input[name='_pass']")
        if password and password.strip():
            self._safe_click(pwd)
            pwd.clear()
            pwd.send_keys(password)
            val = self.driver.execute_script("return arguments[0].value;", pwd)
            if not val:
                self.driver.execute_script("""
                    arguments[0].value = arguments[1];
                    arguments[0].dispatchEvent(new Event('input', {bubbles: true}));
                    arguments[0].dispatchEvent(new Event('change', {bubbles: true}));
                """, pwd, password)
            time.sleep(0.3)

            submit_buttons = self.driver.find_elements(
                By.CSS_SELECTOR,
                "#rcmloginsubmit, button[type='submit'], input[type='submit'], button.btn-primary, #loginbtn, .formbuttons input"
            )
            clicked = False
            for btn in submit_buttons:
                if btn.is_displayed() and btn.is_enabled():
                    self._safe_click(btn)
                    clicked = True
                    break
            if not clicked:
                pwd.send_keys(Keys.ENTER)
        else:
            self.log("No password provided — please enter it in the browser window.", "warning")

        self.log("Waiting for mailbox to authenticate and load...")
        deadline = time.time() + 60
        while time.time() < deadline:
            if self.stop_requested.is_set():
                raise InterruptedError("Stopped during login.")
            if self._any_visible(COMPOSE_CSS, COMPOSE_XPATH) or self.driver.find_elements(
                By.CSS_SELECTOR, "#mailboxlist, #messagelist, #layout, .folderlist"
            ):
                self.log("Webmail authenticated successfully!", "success")
                return
            err_elements = self.driver.find_elements(
                By.CSS_SELECTOR, "#messagestack div, #login-alert, .alert-danger, .boxerror"
            )
            for err in err_elements:
                if err.is_displayed() and err.text.strip():
                    raise Exception(f"Login failed: {err.text.strip()}")
            time.sleep(1)

        raise Exception("Timed out waiting for webmail inbox to load.")

    def _open_compose(self, wait, main_window: str, compose_url: str) -> str:
        before = set(self.driver.window_handles)
        clicked = False
        try:
            self._safe_click(self._visible_one(COMPOSE_CSS, COMPOSE_XPATH, timeout=12))
            clicked = True
        except Exception as exc:
            try:
                self.driver.execute_script("if (window.rcmail) rcmail.command('compose','',null,null);")
                clicked = True
            except Exception:
                pass

        deadline = time.time() + 25
        while time.time() < deadline:
            if self.stop_requested.is_set():
                raise InterruptedError("Stopped waiting for compose window.")
            new_handles = [h for h in self.driver.window_handles if h not in before]
            for handle in new_handles + [main_window]:
                try:
                    self.driver.switch_to.window(handle)
                except Exception:
                    continue
                if self._any_visible(TO_CSS, TO_XPATH):
                    return handle
            time.sleep(0.4)

        # Fallback to compose URL in new tab
        self.driver.switch_to.new_window("tab")
        self.driver.get(compose_url)
        WebDriverWait(self.driver, 20).until(lambda d: self._any_visible(TO_CSS, TO_XPATH))
        return self.driver.current_window_handle

    def _fill_to(self, to_email: str):
        field = self._visible_one(TO_CSS, TO_XPATH, timeout=15)
        self._safe_click(field)
        try:
            field.clear()
        except Exception:
            pass
        field.send_keys(to_email)
        time.sleep(0.3)
        field.send_keys(Keys.ENTER)
        for event in ("input", "change", "blur"):
            self.driver.execute_script(
                "arguments[0].dispatchEvent(new Event(arguments[1],{bubbles:true}));",
                field, event,
            )

    def _fill_subject(self, subject: str):
        field = self._visible_one(SUBJECT_CSS, SUBJECT_XPATH, timeout=12)
        self._safe_click(field)
        try:
            field.clear()
        except Exception:
            pass
        field.send_keys(subject)
        if not (field.get_attribute("value") or "").strip():
            self.driver.execute_script(
                "arguments[0].value=arguments[1];"
                "arguments[0].dispatchEvent(new Event('input',{bubbles:true}));",
                field, subject,
            )

    def _fill_body(self, body_data: Dict[str, str]):
        html_body = body_data.get("html", "")
        plain_body = body_data.get("plain", "")

        # Toggle editor to HTML if available
        try:
            self.driver.execute_script("""
                if (typeof rcmail !== 'undefined') {
                    try {
                        rcmail.command('toggle-editor', {id: 'composebody', html: true});
                    } catch(e) {}
                    var sel = document.querySelector('select[name="_is_html"], select#editor-selector, select[name="_editor_selector"]');
                    if (sel) {
                        for (var i = 0; i < sel.options.length; i++) {
                            if (sel.options[i].value === '1' || sel.options[i].value === 'html') {
                                if (sel.selectedIndex !== i) {
                                    sel.selectedIndex = i;
                                    sel.dispatchEvent(new Event('change', {bubbles: true}));
                                }
                                break;
                            }
                        }
                    }
                }
            """)
        except Exception:
            pass

        self._dismiss_any_dialogs()

        deadline = time.time() + 15
        tinymce_ready = False
        while time.time() < deadline:
            try:
                ready = self.driver.execute_script("""
                    return (typeof tinymce !== 'undefined'
                         && tinymce.get('composebody') !== null
                         && tinymce.get('composebody').initialized);
                """)
                if ready:
                    tinymce_ready = True
                    break
            except Exception:
                pass
            time.sleep(0.4)

        if tinymce_ready:
            try:
                self.driver.execute_script("""
                    var editor = tinymce.get('composebody');
                    editor.setContent(arguments[0]);
                    editor.setDirty(true);
                    if (editor.fire) { editor.fire('change'); }
                    editor.save();
                    var ta = document.getElementById('composebody');
                    if (ta) {
                        ta.value = editor.getContent();
                        ta.dispatchEvent(new Event('input', {bubbles: true}));
                        ta.dispatchEvent(new Event('change', {bubbles: true}));
                    }
                """, html_body)
                time.sleep(0.8)
                return
            except Exception:
                pass

        # Plain text fallback
        try:
            self.driver.execute_script("""
                var ta = document.getElementById('composebody')
                      || document.querySelector("textarea[name='_message']");
                if (ta) {
                    ta.value = arguments[0];
                    ta.dispatchEvent(new Event('input', {bubbles: true}));
                    ta.dispatchEvent(new Event('change', {bubbles: true}));
                }
            """, plain_body)
        except Exception:
            pass

    def _set_priority_high(self):
        try:
            self.driver.execute_script("""
                var sel = document.querySelector('#compose-priority')
                       || document.querySelector("select[name='_priority']");
                if (sel) {
                    sel.value = '1';
                    sel.dispatchEvent(new Event('change', {bubbles: true}));
                }
            """)
        except Exception:
            pass

    def _click_send(self):
        try:
            self._safe_click(self._visible_one(SEND_CSS, SEND_XPATH, timeout=12))
        except Exception:
            self.driver.find_element(By.TAG_NAME, "body").send_keys(Keys.CONTROL, Keys.ENTER)

    def _confirm_sent(self, timeout=20) -> bool:
        deadline = time.time() + timeout
        while time.time() < deadline:
            if self.stop_requested.is_set():
                raise InterruptedError("Stopped while waiting for send confirmation.")
            try:
                for el in self.driver.find_elements(By.CSS_SELECTOR, "#messagestack div, .ui.alert"):
                    text = (el.text or "").lower()
                    if "sent" in text:
                        return True
                    if "error" in text or "failed" in text or "invalid" in text:
                        raise Exception(f"Webmail reported: {el.text.strip()}")
            except Exception as exc:
                if "Webmail reported" in str(exc):
                    raise
            if not self._any_visible(SUBJECT_CSS, SUBJECT_XPATH):
                return True
            time.sleep(0.4)
        return False

    def _reset_to_inbox(self, main_window: str):
        try:
            for handle in list(self.driver.window_handles):
                if handle != main_window:
                    self.driver.switch_to.window(handle)
                    self.driver.close()
            self.driver.switch_to.window(main_window)
            self._dismiss_any_dialogs()
        except Exception:
            pass

    # -----------------------------------------------------------------------
    # MAIN WORKER RUN LOOP
    # -----------------------------------------------------------------------
    def _run_job(self, config: Dict[str, Any]):
        webmail_url = config.get("webmail_url", "https://us3.webmail.mailhostbox.com").rstrip("/")
        from_email = config.get("from_email", "").strip()
        password = config.get("password", "")
        delay = float(config.get("delay_seconds", 30))
        headless = bool(config.get("headless", False))
        log_file = config.get("log_file", "send_log.csv")
        recipients = config.get("recipients", [])
        template_data = config.get("template_data", {"html": "", "plain": ""})
        custom_subjects = config.get("custom_subjects", [])
        compose_url = webmail_url + "/?_task=mail&_action=compose"

        self.total_count = len(recipients)
        self.log(f"Starting bulk mailer worker for {self.total_count} recipients.")
        self.log(f"From: {from_email} | Delay: {delay}s | Headless: {headless}")

        try:
            self.driver = self._start_browser(headless=headless)
            wait = WebDriverWait(self.driver, 25)
            self._login(wait, webmail_url, from_email, password)
            main_window = self.driver.current_window_handle

            for i, r in enumerate(recipients):
                # Check for stop
                if self.stop_requested.is_set():
                    self.log("Worker stopped by user.", "warning")
                    break

                # Check for pause
                while self.pause_requested.is_set():
                    self.status = "paused"
                    time.sleep(1)
                    if self.stop_requested.is_set():
                        break

                if self.stop_requested.is_set():
                    break

                self.status = "running"
                self.current_index = i + 1
                self.current_recipient = f"{r['name']} <{r['email']}>"

                if custom_subjects:
                    subject = random.choice(custom_subjects)
                else:
                    subject = generate_fallback_subject(r.get("name"), r.get("topic"), i)

                self.current_subject = subject
                body = personalize_template(template_data, r.get("name", ""), r.get("paper_title", ""))

                self.log(f"[{i+1}/{self.total_count}] Dispatching to {r['email']} with subject: '{subject}'...")

                try:
                    self._dismiss_any_dialogs()
                    compose_window = self._open_compose(wait, main_window, compose_url)
                    self.driver.switch_to.window(compose_window)

                    self._fill_to(r["email"])
                    self._fill_subject(subject)
                    self._fill_body(body)
                    self._set_priority_high()
                    self._click_send()

                    if not self._confirm_sent():
                        raise Exception("Compose window did not close after send.")

                    if compose_window != main_window and compose_window in self.driver.window_handles:
                        try:
                            self.driver.close()
                        except Exception:
                            pass
                    self.driver.switch_to.window(main_window)
                    self._dismiss_any_dialogs()

                    # Record success
                    self.sent_count += 1
                    self._log_to_file(log_file, r["email"], r["name"], subject, "sent")
                    self.log(f"[{i+1}/{self.total_count}] SUCCESS -> {r['email']}", "success")
                    self.emit_progress(email=r["email"], status_code="sent", subject=subject)

                except Exception as exc:
                    err_str = str(exc)
                    # Check if browser was closed or disconnected
                    is_fatal = any(p in err_str.lower() for p in [
                        "httpconnectionpool", "connection refused", "actively refused",
                        "max retries exceeded", "newconnectionerror", "invalid session id",
                        "no such window", "target window already closed", "session deleted",
                        "chrome not reachable", "disconnected"
                    ])

                    if is_fatal or not self._is_driver_alive():
                        self.log("[CRITICAL] Chrome browser was closed or lost connection. Stopping campaign.", "error")
                        self._log_to_file(log_file, r["email"], r["name"], subject, "failed", "Chrome closed or disconnected")
                        self.status = "error"
                        self.error_message = "Chrome browser was closed or lost connection."
                        self.failed_count += 1
                        self.emit_progress(email=r["email"], status_code="failed", subject=subject, error="Chrome closed or disconnected")
                        break

                    self._dump_error(r["email"].replace("@", "_at_"))
                    self.failed_count += 1
                    self._log_to_file(log_file, r["email"], r["name"], subject, "failed", err_str)
                    self.log(f"[{i+1}/{self.total_count}] FAILED -> {r['email']}: {err_str}", "error")
                    self.emit_progress(email=r["email"], status_code="failed", subject=subject, error=err_str)
                    self._reset_to_inbox(main_window)

                # Delay before next email
                if i < self.total_count - 1 and not self.stop_requested.is_set():
                    self.log(f"Waiting {delay} seconds before next recipient...")
                    deadline = time.time() + delay
                    while time.time() < deadline:
                        if self.stop_requested.is_set():
                            break
                        time.sleep(0.5)

            if not self.stop_requested.is_set():
                if self.status != "error":
                    self.status = "completed"
                    self.log(f"Process finished! Sent: {self.sent_count}, Failed: {self.failed_count} of {self.total_count}.", "success")
            else:
                self.status = "stopped"

        except Exception as exc:
            self.status = "error"
            self.error_message = str(exc)
            self.log(f"Critical execution error: {exc}", "error")
        finally:
            if self.driver:
                try:
                    self.driver.quit()
                except Exception:
                    pass
                self.driver = None
            self.emit_progress()

    def _is_driver_alive(self) -> bool:
        if not self.driver:
            return False
        try:
            _ = self.driver.window_handles
            return True
        except Exception:
            return False

    def _log_to_file(self, log_path: str, email: str, name: str, subject: str, status: str, error: str = ""):
        exists = Path(log_path).exists()
        try:
            with open(log_path, "a", newline="", encoding="utf-8", errors="replace") as f:
                w = csv.writer(f)
                if not exists:
                    w.writerow(["timestamp", "email", "name", "subject", "status", "error"])
                w.writerow([time.strftime("%Y-%m-%d %H:%M:%S"), email, name, subject, status, error])
        except PermissionError:
            # Fallback if send_log.csv is open in Excel
            fb = Path("send_log_pending.csv")
            fb_exists = fb.exists()
            try:
                with open(fb, "a", newline="", encoding="utf-8", errors="replace") as f:
                    w = csv.writer(f)
                    if not fb_exists:
                        w.writerow(["timestamp", "email", "name", "subject", "status", "error"])
                    w.writerow([time.strftime("%Y-%m-%d %H:%M:%S"), email, name, subject, status, error])
                self.log("send_log.csv is open in Excel! Log written to send_log_pending.csv.", "warning")
            except Exception:
                pass
